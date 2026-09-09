import csv
import io
import json
import os
import re
import uuid
import zipfile
from datetime import datetime, timezone
from functools import wraps
from urllib.parse import urlparse

import boto3
import pg8000.dbapi as pg8000
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from flask import (
    Flask, render_template, request, redirect, url_for,
    flash, jsonify, session, Response, abort
)
from PIL import Image, ImageDraw, ImageFont
from io import BytesIO

try:  # optional — only used to draw Arabic text on the shareable card
    import arabic_reshaper
    from bidi.algorithm import get_display
except Exception:  # pragma: no cover
    arabic_reshaper = None
    get_display = None

load_dotenv()

# ---------------------------------------------------------------------------
# Configuration (all from environment variables — see .env.example)
# ---------------------------------------------------------------------------
DATABASE_URL = os.environ.get('DATABASE_URL', '')

R2_ACCOUNT_ID = os.environ.get('R2_ACCOUNT_ID', '')
R2_ACCESS_KEY_ID = os.environ.get('R2_ACCESS_KEY_ID', '')
R2_SECRET_ACCESS_KEY = os.environ.get('R2_SECRET_ACCESS_KEY', '')
R2_BUCKET_NAME = os.environ.get('R2_BUCKET_NAME', '')
R2_PUBLIC_URL = os.environ.get('R2_PUBLIC_URL', '').rstrip('/')

ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'changeme123')
VIEWER_PASSWORD = os.environ.get('VIEWER_PASSWORD', 'change-me')
SECRET_KEY = os.environ.get('SECRET_KEY', 'change-this-secret-key-before-deploying')

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'webp'}

app = Flask(__name__)
app.secret_key = SECRET_KEY
# Generous cap: a product can now be uploaded with several photos in one go.
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024


# ---------------------------------------------------------------------------
# Database (Postgres / Supabase via pg8000 — pure Python, no compiler needed)
# ---------------------------------------------------------------------------
def get_db():
    if not DATABASE_URL:
        raise RuntimeError(
            'DATABASE_URL is not set. Add it to your .env file (see .env.example).'
        )
    parsed = urlparse(DATABASE_URL)
    return pg8000.connect(
        user=parsed.username,
        password=parsed.password,
        host=parsed.hostname,
        port=parsed.port or 5432,
        database=parsed.path.lstrip('/'),
        ssl_context=True,
    )


def init_db():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        CREATE TABLE IF NOT EXISTS items (
            id TEXT PRIMARY KEY,
            name TEXT,
            category TEXT,
            image_key TEXT NOT NULL,
            thumb_key TEXT NOT NULL,
            date_added TIMESTAMP NOT NULL
        )
    ''')
    # One product can have many photos. items.image_key / thumb_key still hold
    # the "cover" photo (the first one) so the catalog grid needs no changes.
    cur.execute('''
        CREATE TABLE IF NOT EXISTS item_images (
            id TEXT PRIMARY KEY,
            item_id TEXT NOT NULL REFERENCES items(id) ON DELETE CASCADE,
            image_key TEXT NOT NULL,
            thumb_key TEXT NOT NULL,
            position INTEGER NOT NULL DEFAULT 0,
            date_added TIMESTAMP NOT NULL
        )
    ''')
    cur.execute('CREATE INDEX IF NOT EXISTS idx_item_images_item ON item_images (item_id)')
    # Backfill: give every existing product its current photo as the first image.
    cur.execute('''
        INSERT INTO item_images (id, item_id, image_key, thumb_key, position, date_added)
        SELECT 'legacy-' || i.id, i.id, i.image_key, i.thumb_key, 0, i.date_added
        FROM items i
        WHERE NOT EXISTS (SELECT 1 FROM item_images im WHERE im.item_id = i.id)
        ON CONFLICT (id) DO NOTHING
    ''')
    conn.commit()
    cur.close()
    conn.close()


def rows_to_dicts(cur, rows):
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, row)) for row in rows]


def query_items(q):
    conn = get_db()
    cur = conn.cursor()
    if q:
        like = f'%{q}%'
        cur.execute(
            'SELECT * FROM items WHERE id LIKE %s OR name LIKE %s OR category LIKE %s ORDER BY date_added DESC',
            (like, like, like)
        )
    else:
        cur.execute('SELECT * FROM items ORDER BY date_added DESC')
    items = rows_to_dicts(cur, cur.fetchall())
    cur.close()
    conn.close()
    return items


def get_item(item_id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM items WHERE id = %s', (item_id,))
    rows = cur.fetchall()
    items = rows_to_dicts(cur, rows)
    cur.close()
    conn.close()
    return items[0] if items else None


def get_item_images(item_id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT * FROM item_images WHERE item_id = %s ORDER BY position, date_added',
        (item_id,)
    )
    images = rows_to_dicts(cur, cur.fetchall())
    cur.close()
    conn.close()
    return images


def attach_thumbs(items):
    """Attach each item's ordered thumbnail URLs as item['thumb_urls'] (for the
    catalog hover preview). Falls back to the cover thumb if none are found."""
    if not items:
        return items
    ids = [it['id'] for it in items]
    placeholders = ','.join(['%s'] * len(ids))
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        f'SELECT item_id, thumb_key FROM item_images '
        f'WHERE item_id IN ({placeholders}) ORDER BY position, date_added',
        ids
    )
    by_item = {}
    for item_id, thumb_key in cur.fetchall():
        by_item.setdefault(item_id, []).append(r2_url(thumb_key))
    cur.close()
    conn.close()
    for it in items:
        it['thumb_urls'] = by_item.get(it['id']) or [r2_url(it['thumb_key'])]
    return items


def all_items_with_images():
    """Every item plus its ordered images — used by the CSV / JSON export."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        'SELECT id, name, category, image_key, thumb_key, date_added '
        'FROM items ORDER BY date_added DESC'
    )
    items = rows_to_dicts(cur, cur.fetchall())
    cur.execute(
        'SELECT item_id, position, image_key, thumb_key FROM item_images '
        'ORDER BY item_id, position, date_added'
    )
    by_item = {}
    for item_id, position, image_key, thumb_key in cur.fetchall():
        by_item.setdefault(item_id, []).append({
            'position': position,
            'image_url': r2_url(image_key),
            'thumb_url': r2_url(thumb_key),
        })
    cur.close()
    conn.close()
    for it in items:
        it['images'] = by_item.get(it['id'], [])
    return items


def all_items_with_image_keys():
    """Every item plus its ordered R2 image keys — used by the ZIP export."""
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT id, name, category, date_added FROM items ORDER BY date_added DESC')
    items = rows_to_dicts(cur, cur.fetchall())
    cur.execute(
        'SELECT item_id, image_key FROM item_images ORDER BY item_id, position, date_added'
    )
    by_item = {}
    for item_id, image_key in cur.fetchall():
        by_item.setdefault(item_id, []).append(image_key)
    cur.close()
    conn.close()
    for it in items:
        it['image_keys'] = by_item.get(it['id'], [])
    return items


# ---------------------------------------------------------------------------
# Storage (Cloudflare R2, via the S3-compatible API)
# ---------------------------------------------------------------------------
def get_r2_client():
    if not (R2_ACCOUNT_ID and R2_ACCESS_KEY_ID and R2_SECRET_ACCESS_KEY):
        raise RuntimeError(
            'R2 credentials are not set. Add them to your .env file (see .env.example).'
        )
    return boto3.client(
        's3',
        endpoint_url=f'https://{R2_ACCOUNT_ID}.r2.cloudflarestorage.com',
        aws_access_key_id=R2_ACCESS_KEY_ID,
        aws_secret_access_key=R2_SECRET_ACCESS_KEY,
        region_name='auto',
    )


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def make_thumbnail_bytes(file_bytes, size=(500, 500)):
    with Image.open(BytesIO(file_bytes)) as img:
        img = img.convert('RGB')
        img.thumbnail(size)
        out = BytesIO()
        img.save(out, 'JPEG', quality=85)
        return out.getvalue()


def save_image(file):
    """Upload an image + thumbnail to R2. Returns (image_key, thumb_key)."""
    ext = file.filename.rsplit('.', 1)[1].lower()
    file_bytes = file.read()
    unique_name = uuid.uuid4().hex

    image_key = f'items/{unique_name}.{ext}'
    thumb_key = f'items/thumbs/{unique_name}.jpg'

    content_type = file.mimetype or f'image/{ext}'
    s3 = get_r2_client()
    s3.put_object(Bucket=R2_BUCKET_NAME, Key=image_key, Body=file_bytes, ContentType=content_type)

    try:
        thumb_bytes = make_thumbnail_bytes(file_bytes)
        s3.put_object(Bucket=R2_BUCKET_NAME, Key=thumb_key, Body=thumb_bytes, ContentType='image/jpeg')
    except Exception:
        thumb_key = image_key  # fallback: reuse original if thumbnailing fails

    return image_key, thumb_key


def delete_keys(image_key, thumb_key):
    """Best-effort removal of one image + its thumbnail from R2."""
    try:
        s3 = get_r2_client()
        for key in {image_key, thumb_key}:
            if key:
                s3.delete_object(Bucket=R2_BUCKET_NAME, Key=key)
    except Exception:
        pass  # don't block the DB change if R2 cleanup fails


def delete_all_item_images(images):
    for img in images:
        delete_keys(img['image_key'], img['thumb_key'])


def r2_url(key):
    return f'{R2_PUBLIC_URL}/{key}'


def r2_get_bytes(key):
    """Download one object's bytes from R2 (server-side, no CORS involved)."""
    s3 = get_r2_client()
    obj = s3.get_object(Bucket=R2_BUCKET_NAME, Key=key)
    return obj['Body'].read()


# ---------------------------------------------------------------------------
# Shareable product card (server-side PNG, drawn with Pillow)
# ---------------------------------------------------------------------------
CARD_W, CARD_H, CARD_PHOTO_H = 1080, 1350, 1040
CARD_MARGIN = 48
FONT_PATH = os.path.join(os.path.dirname(__file__), 'static', 'fonts',
                         'NotoNaskhArabic-VariableFont.ttf')
_ARABIC_RE = re.compile('[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]')


def _card_font(size, weight=400):
    try:
        font = ImageFont.truetype(FONT_PATH, size)
        try:
            font.set_variation_by_axes([weight])
        except Exception:
            pass
        return font
    except Exception:  # font file missing — degrade rather than crash
        try:
            return ImageFont.load_default(size=size)
        except TypeError:
            return ImageFont.load_default()


def _shape(text):
    """Return (display_text, is_rtl). Arabic is reshaped + bidi-ordered so it
    renders correctly without a complex-text layout engine."""
    if not text or not _ARABIC_RE.search(text):
        return text, False
    if arabic_reshaper and get_display:
        try:
            return get_display(arabic_reshaper.reshape(text), base_dir='R'), True
        except Exception:
            return text, True
    return text, True


def _cover_crop(img, target_w, target_h):
    iw, ih = img.size
    scale = max(target_w / iw, target_h / ih)
    nw, nh = max(1, int(iw * scale)), max(1, int(ih * scale))
    img = img.resize((nw, nh), Image.LANCZOS)
    left, top = (nw - target_w) // 2, (nh - target_h) // 2
    return img.crop((left, top, left + target_w, top + target_h))


def _ellipsize(text, limit):
    text = ' '.join(text.split())
    return text if len(text) <= limit else text[:limit - 1].rstrip() + '…'


def _draw_line(draw, text, y, size, colour, weight=400):
    disp, rtl = _shape(_ellipsize(text, 44))
    if rtl:
        draw.text((CARD_W - CARD_MARGIN, y), disp, font=_card_font(size, weight),
                  fill=colour, anchor='ra')
    else:
        draw.text((CARD_MARGIN, y), disp, font=_card_font(size, weight),
                  fill=colour, anchor='la')


def compose_card_png(item, photo_bytes):
    card = Image.new('RGB', (CARD_W, CARD_H), '#FFFFFF')
    try:
        with Image.open(BytesIO(photo_bytes)) as p:
            card.paste(_cover_crop(p.convert('RGB'), CARD_W, CARD_PHOTO_H), (0, 0))
    except Exception:
        pass

    d = ImageDraw.Draw(card)
    d.line([(0, CARD_PHOTO_H), (CARD_W, CARD_PHOTO_H)], fill='#DAD5C9', width=2)

    y = CARD_PHOTO_H + 44
    d.text((CARD_MARGIN, y), f"ID {item['id']}", font=_card_font(66, 700), fill='#232019')
    y += 96
    name = (item.get('name') or '').strip()
    if name:
        _draw_line(d, name, y, 42, '#232019', 600)
        y += 58
    category = (item.get('category') or '').strip()
    if category:
        _draw_line(d, category, y, 32, '#6E7350')

    d.text((CARD_W - CARD_MARGIN, CARD_H - 40), 'Rack & ID', font=_card_font(26),
           fill='#8A8477', anchor='rs')

    out = BytesIO()
    card.save(out, 'PNG')
    return out.getvalue()


def safe_filename(value):
    return re.sub(r'[^A-Za-z0-9._-]+', '_', value).strip('_') or 'item'


# ---------------------------------------------------------------------------
# Translations (basic English / Arabic)
# ---------------------------------------------------------------------------
TRANSLATIONS = {
    'en': {
        'app_name': 'Rack & ID',
        'nav_catalog': 'Catalog',
        'nav_add_item': 'Add item',
        'nav_manage': 'Manage',
        'nav_login': 'Sign in',
        'nav_logout': 'Log out',
        'search_label': 'Find by ID or name',
        'search_placeholder': 'e.g. 4521 or "denim jacket"',
        'search_btn': 'Search',
        'clear_link': 'Clear',
        'items_count': 'items in catalog',
        'item_count_singular': 'item in catalog',
        'empty_no_match': 'No items match',
        'clear_search': 'Clear search',
        'empty_no_items': 'No items yet. Start by adding the first one.',
        'add_item_cta': 'Add an item →',
        'untitled_item': 'Untitled item',
        'remove_btn': 'Remove',
        'edit_btn': 'Edit',
        'view_full': 'View full size',
        'back_to_catalog': '← Back to catalog',
        'item_id_label': 'Item ID',
        'category_label': 'Category',
        'added_label': 'Added',
        'print_btn': 'Print',
        'save_image_btn': 'Save as image',
        'manage_title': 'Manage catalog',
        'manage_items': 'Items',
        'manage_import_hint': 'Restore from a JSON or ZIP export.',
        'export_label': 'Export catalog',
        'export_zip_label': 'ZIP + photos',
        'import_link': 'Import / restore',
        'import_title': 'Import / restore',
        'import_sub': 'Upload a JSON or ZIP export. Items whose ID is already in the catalog are skipped — nothing existing is changed.',
        'import_field': 'Export file (.json or .zip)',
        'import_btn': 'Import',
        'import_latest_btn': 'Restore the latest automatic backup',
        'import_no_file': 'Choose a file to import.',
        'import_bad_type': 'Upload a .json or .zip export file.',
        'import_failed': 'Import failed.',
        'import_done': 'Imported {added} item(s) and {photos} photo(s). Skipped {skipped} already in the catalog.',
        'import_nolatest': 'No automatic backup found yet.',
        'photos_word': 'photos',
        'item_not_found': 'Item not found.',
        'upload_title': 'Add a new item',
        'upload_sub': 'Use the same ID as your existing system so both stay in sync.',
        'field_id': 'Item ID',
        'field_id_hint': '(required, must match your system)',
        'field_name': 'Name',
        'field_optional': '(optional)',
        'field_category': 'Category',
        'field_photo': 'Photo',
        'field_photos': 'Photos',
        'field_photos_hint': '(one or more)',
        'no_photo_selected': 'No photo selected',
        'no_photos_selected': 'No photos selected',
        'edit_photos_current': 'Current photos',
        'edit_photos_add': 'Add more photos',
        'remove_photo': 'Remove',
        'submit_add': 'Add to catalog',
        'edit_title': 'Edit item',
        'edit_sub': 'Tick any photo to remove it, and add more below if you like — the rest stay as they are.',
        'submit_edit': 'Save changes',
        'login_title': 'Sign in',
        'login_sub': 'Enter the viewer password to browse the catalog, or the admin password to add, edit, or remove items.',
        'field_password': 'Password',
        'login_btn': 'Log in',
        'login_required_msg': 'Please log in as admin to do that.',
        'login_required_view': 'Please log in to view the catalog.',
        'login_success': 'Logged in as admin.',
        'login_success_viewer': 'Logged in.',
        'login_failed': 'Incorrect password.',
        'logout_success': 'Logged out.',
        'viewer_note': 'Viewing as guest — log in as admin to add or edit items.',
        'id_exists': 'An item with this ID already exists. Use a different ID.',
        'id_required': 'Item ID is required.',
        'image_required': 'At least one photo is required.',
        'bad_filetype': 'Only PNG, JPG, and WEBP images are supported.',
        'item_added': 'Added to the catalog.',
        'item_updated': 'Changes saved.',
    },
    'ar': {
        'app_name': 'رفّ آي دي',
        'nav_catalog': 'الكتالوج',
        'nav_add_item': 'إضافة عنصر',
        'nav_manage': 'الإدارة',
        'nav_login': 'تسجيل الدخول',
        'nav_logout': 'تسجيل الخروج',
        'search_label': 'ابحث بالرقم التعريفي أو الاسم',
        'search_placeholder': 'مثال: 4521 أو "جاكيت جينز"',
        'search_btn': 'بحث',
        'clear_link': 'مسح',
        'items_count': 'عنصر في الكتالوج',
        'item_count_singular': 'عنصر واحد في الكتالوج',
        'empty_no_match': 'لا توجد عناصر مطابقة لـ',
        'clear_search': 'مسح البحث',
        'empty_no_items': 'لا توجد عناصر بعد. ابدأ بإضافة أول عنصر.',
        'add_item_cta': 'أضف عنصرًا ←',
        'untitled_item': 'عنصر بدون اسم',
        'remove_btn': 'حذف',
        'edit_btn': 'تعديل',
        'view_full': 'عرض بالحجم الكامل',
        'back_to_catalog': '→ العودة إلى الكتالوج',
        'item_id_label': 'الرقم التعريفي',
        'category_label': 'الفئة',
        'added_label': 'تمت الإضافة',
        'print_btn': 'طباعة',
        'save_image_btn': 'حفظ كصورة',
        'manage_title': 'إدارة الكتالوج',
        'manage_items': 'العناصر',
        'manage_import_hint': 'استعادة من ملف تصدير JSON أو ZIP.',
        'export_label': 'تصدير الكتالوج',
        'export_zip_label': 'ZIP مع الصور',
        'import_link': 'استيراد / استعادة',
        'import_title': 'استيراد / استعادة',
        'import_sub': 'ارفع ملف تصدير بصيغة JSON أو ZIP. يتم تجاهل العناصر التي رقمها التعريفي موجود مسبقًا — ولا يتغيّر أي شيء قائم.',
        'import_field': 'ملف التصدير (.json أو .zip)',
        'import_btn': 'استيراد',
        'import_latest_btn': 'استعادة آخر نسخة احتياطية تلقائية',
        'import_no_file': 'اختر ملفًا للاستيراد.',
        'import_bad_type': 'ارفع ملف تصدير بصيغة .json أو .zip.',
        'import_failed': 'فشل الاستيراد.',
        'import_done': 'تم استيراد {added} عنصرًا و{photos} صورة. وتم تجاهل {skipped} موجودًا مسبقًا.',
        'import_nolatest': 'لا توجد نسخة احتياطية تلقائية بعد.',
        'photos_word': 'صور',
        'item_not_found': 'العنصر غير موجود.',
        'upload_title': 'إضافة عنصر جديد',
        'upload_sub': 'استخدم نفس الرقم التعريفي في نظامك الحالي ليبقى الاثنان متطابقين.',
        'field_id': 'الرقم التعريفي للعنصر',
        'field_id_hint': '(مطلوب، يجب أن يطابق نظامك)',
        'field_name': 'الاسم',
        'field_optional': '(اختياري)',
        'field_category': 'الفئة',
        'field_photo': 'الصورة',
        'field_photos': 'الصور',
        'field_photos_hint': '(صورة واحدة أو أكثر)',
        'no_photo_selected': 'لم يتم اختيار صورة',
        'no_photos_selected': 'لم يتم اختيار صور',
        'edit_photos_current': 'الصور الحالية',
        'edit_photos_add': 'إضافة المزيد من الصور',
        'remove_photo': 'حذف',
        'submit_add': 'إضافة إلى الكتالوج',
        'edit_title': 'تعديل العنصر',
        'edit_sub': 'حدد أي صورة لحذفها، وأضف المزيد أدناه إن رغبت — وتبقى بقية الصور كما هي.',
        'submit_edit': 'حفظ التغييرات',
        'login_title': 'تسجيل الدخول',
        'login_sub': 'أدخل كلمة مرور العرض لتصفح الكتالوج، أو كلمة مرور المسؤول لإضافة العناصر أو تعديلها أو حذفها.',
        'field_password': 'كلمة المرور',
        'login_btn': 'تسجيل الدخول',
        'login_required_msg': 'الرجاء تسجيل الدخول كمسؤول للقيام بذلك.',
        'login_required_view': 'الرجاء تسجيل الدخول لعرض الكتالوج.',
        'login_success': 'تم تسجيل الدخول كمسؤول.',
        'login_success_viewer': 'تم تسجيل الدخول.',
        'login_failed': 'كلمة المرور غير صحيحة.',
        'logout_success': 'تم تسجيل الخروج.',
        'viewer_note': 'أنت تتصفح كزائر — سجّل الدخول كمسؤول لإضافة العناصر أو تعديلها.',
        'id_exists': 'يوجد عنصر بهذا الرقم التعريفي مسبقًا. استخدم رقمًا مختلفًا.',
        'id_required': 'الرقم التعريفي للعنصر مطلوب.',
        'image_required': 'مطلوب صورة واحدة على الأقل.',
        'bad_filetype': 'يتم دعم صور PNG وJPG وWEBP فقط.',
        'item_added': 'تمت الإضافة إلى الكتالوج.',
        'item_updated': 'تم حفظ التغييرات.',
    },
}


DEFAULT_LANG = os.environ.get('DEFAULT_LANG', 'ar')


def t(key):
    lang = session.get('lang', DEFAULT_LANG)
    return TRANSLATIONS.get(lang, TRANSLATIONS['en']).get(key, key)


@app.context_processor
def inject_globals():
    lang = session.get('lang', DEFAULT_LANG)
    return dict(
        t=t,
        current_lang=lang,
        is_rtl=(lang == 'ar'),
        is_admin=bool(session.get('is_admin')),
        is_viewer=bool(session.get('is_viewer') or session.get('is_admin')),
        r2_url=r2_url,
    )


@app.route('/lang/<code>')
def set_lang(code):
    if code in TRANSLATIONS:
        session['lang'] = code
    return redirect(request.referrer or url_for('index'))


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('is_admin'):
            flash(t('login_required_msg'), 'error')
            return redirect(url_for('login', next=request.path))
        return f(*args, **kwargs)
    return decorated


def viewer_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not (session.get('is_viewer') or session.get('is_admin')):
            flash(t('login_required_view'), 'error')
            return redirect(url_for('login', next=request.path))
        return f(*args, **kwargs)
    return decorated


@app.route('/login', methods=['GET', 'POST'])
def login():
    next_url = request.values.get('next') or url_for('index')
    if request.method == 'POST':
        password = request.form.get('password', '')
        if ADMIN_PASSWORD and password == ADMIN_PASSWORD:
            session['is_admin'] = True
            session['is_viewer'] = True
            flash(t('login_success'), 'success')
            return redirect(request.form.get('next') or url_for('index'))
        if VIEWER_PASSWORD and password == VIEWER_PASSWORD:
            session['is_viewer'] = True
            session.pop('is_admin', None)
            flash(t('login_success_viewer'), 'success')
            return redirect(request.form.get('next') or url_for('index'))
        flash(t('login_failed'), 'error')
    return render_template('login.html', next=next_url)


@app.route('/logout', methods=['POST'])
def logout():
    session.pop('is_admin', None)
    session.pop('is_viewer', None)
    flash(t('logout_success'), 'success')
    return redirect(url_for('login'))


# ---------------------------------------------------------------------------
# Routes — viewing (any signed-in role: viewer or admin)
# ---------------------------------------------------------------------------
@app.route('/')
@viewer_required
def index():
    q = request.args.get('q', '').strip()
    items = attach_thumbs(query_items(q))
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) FROM items')
    total = cur.fetchone()[0]
    cur.close()
    conn.close()
    return render_template('index.html', items=items, query=q, total=total)


@app.route('/api/search')
@viewer_required
def api_search():
    q = request.args.get('q', '').strip()
    items = query_items(q)
    for item in items:
        if isinstance(item.get('date_added'), datetime):
            item['date_added'] = item['date_added'].isoformat()
    return jsonify(items)


@app.route('/item/<item_id>')
@viewer_required
def item_detail(item_id):
    item = get_item(item_id)
    if not item:
        flash(t('item_not_found'), 'error')
        return redirect(url_for('index'))
    return render_template('item_detail.html', item=item, images=get_item_images(item_id))


@app.route('/item/<item_id>/card.png')
@viewer_required
def item_card(item_id):
    item = get_item(item_id)
    if not item:
        abort(404)
    images = get_item_images(item_id)
    keys = [im['image_key'] for im in images] or [item['image_key']]
    try:
        idx = int(request.args.get('img', 0))
    except (TypeError, ValueError):
        idx = 0
    idx = max(0, min(idx, len(keys) - 1))
    try:
        photo_bytes = r2_get_bytes(keys[idx])
    except Exception:
        photo_bytes = b''
    png = compose_card_png(item, photo_bytes)
    return Response(png, mimetype='image/png', headers={
        'Content-Disposition': f'attachment; filename="item-{safe_filename(item_id)}.png"',
        'Cache-Control': 'no-store',
    })


# ---------------------------------------------------------------------------
# Routes — full catalog export (admin only)
# ---------------------------------------------------------------------------
@app.route('/export/items.csv')
@login_required
def export_csv():
    items = all_items_with_images()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(['id', 'name', 'category', 'date_added', 'photo_count', 'photo_urls'])
    for it in items:
        added = it['date_added'].isoformat() if it['date_added'] else ''
        writer.writerow([
            it['id'], it['name'] or '', it['category'] or '', added,
            len(it['images']), ' | '.join(i['image_url'] for i in it['images']),
        ])
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    # UTF-8 BOM so Excel reads Arabic names correctly.
    body = '﻿' + buf.getvalue()
    return Response(body, mimetype='text/csv; charset=utf-8', headers={
        'Content-Disposition': f'attachment; filename="rack-and-id-{stamp}.csv"',
    })


def build_backup_payload():
    """The JSON snapshot used by both the export button and backup.py."""
    items = all_items_with_images()
    return {
        'exported_at': datetime.now(timezone.utc).isoformat(),
        'count': len(items),
        'r2_public_url': R2_PUBLIC_URL,
        'items': [{
            'id': it['id'],
            'name': it['name'],
            'category': it['category'],
            'date_added': it['date_added'].isoformat() if it['date_added'] else None,
            'images': it['images'],
        } for it in items],
    }


@app.route('/export/items.json')
@login_required
def export_json():
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    body = json.dumps(build_backup_payload(), ensure_ascii=False, indent=2)
    return Response(body, mimetype='application/json; charset=utf-8', headers={
        'Content-Disposition': f'attachment; filename="rack-and-id-{stamp}.json"',
    })


@app.route('/export/items.xlsx')
@login_required
def export_xlsx():
    items = all_items_with_images()
    max_photos = min(10, max((len(it['images']) for it in items), default=0))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Catalog'
    ws.sheet_view.rightToLeft = (session.get('lang', DEFAULT_LANG) == 'ar')

    headers = ['ID', 'Name', 'Category', 'Date added', 'Photos']
    headers += [f'Photo {i}' for i in range(1, max_photos + 1)]
    ws.append(headers)
    head_fill = PatternFill('solid', fgColor='33455E')
    head_font = Font(bold=True, color='FFFFFF')
    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'

    link_font = Font(color='0563C1', underline='single')
    for it in items:
        added = it['date_added'].strftime('%Y-%m-%d') if it['date_added'] else ''
        ws.append([it['id'], it['name'] or '', it['category'] or '', added, len(it['images'])])
        row = ws.max_row
        for i, img in enumerate(it['images'][:max_photos]):
            cell = ws.cell(row=row, column=6 + i, value=f'Photo {i + 1}')
            cell.hyperlink = img['image_url']
            cell.font = link_font

    for i, width in enumerate([16, 34, 22, 14, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for i in range(6, 6 + max_photos):
        ws.column_dimensions[get_column_letter(i)].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    return Response(buf.getvalue(), headers={
        'Content-Disposition': f'attachment; filename="rack-and-id-{stamp}.xlsx"',
    }, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/export/catalog.zip')
@login_required
def export_zip():
    """CSV plus every photo file, named by item ID. Downloads each image from
    R2 in turn, so this is an occasional-use admin action, not a hot path."""
    items = all_items_with_image_keys()
    s3 = get_r2_client()

    index = io.StringIO()
    writer = csv.writer(index)
    writer.writerow(['id', 'name', 'category', 'date_added', 'photo_count', 'photo_files'])

    mem = io.BytesIO()
    with zipfile.ZipFile(mem, 'w', zipfile.ZIP_DEFLATED) as zf:
        for it in items:
            safe_id = safe_filename(it['id'])
            saved = []
            for i, key in enumerate(it['image_keys'], 1):
                ext = key.rsplit('.', 1)[-1].lower() if '.' in key else 'jpg'
                arcname = f'images/{safe_id}_{i}.{ext}'
                try:
                    body = s3.get_object(Bucket=R2_BUCKET_NAME, Key=key)['Body'].read()
                    zf.writestr(arcname, body)
                    saved.append(arcname.split('/', 1)[1])
                except Exception:
                    pass
            added = it['date_added'].isoformat() if it['date_added'] else ''
            writer.writerow([it['id'], it['name'] or '', it['category'] or '',
                             added, len(it['image_keys']), ' | '.join(saved)])
        zf.writestr('catalog.csv', '﻿' + index.getvalue())

    stamp = datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    return Response(mem.getvalue(), mimetype='application/zip', headers={
        'Content-Disposition': f'attachment; filename="rack-and-id-{stamp}.zip"',
    })


# ---------------------------------------------------------------------------
# Import / restore (admin only) — adds items whose ID is not already present
# ---------------------------------------------------------------------------
def _key_from_url(url):
    if not url:
        return None
    prefix = R2_PUBLIC_URL + '/'
    if R2_PUBLIC_URL and url.startswith(prefix):
        return url[len(prefix):]
    parts = url.split('/', 3)
    return parts[3] if len(parts) == 4 else None


def _parse_dt(value):
    if value:
        try:
            return datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        except ValueError:
            pass
    return datetime.now(timezone.utc)


def _existing_ids(cur):
    cur.execute('SELECT id FROM items')
    return {row[0] for row in cur.fetchall()}


def import_from_json(raw):
    """Restore item + image rows from a JSON export. Image keys are derived
    from the stored URLs, so the R2 objects must still exist."""
    data = json.loads(raw.decode('utf-8'))
    records = data.get('items') if isinstance(data, dict) else data
    records = records or []
    added = photos = skipped = 0
    conn = get_db()
    cur = conn.cursor()
    have = _existing_ids(cur)
    for rec in records:
        iid = str(rec.get('id', '')).strip()
        if not iid or iid in have:
            skipped += 1
            continue
        parsed = []
        for im in (rec.get('images') or []):
            ik = _key_from_url(im.get('image_url'))
            if not ik:
                continue
            tk = _key_from_url(im.get('thumb_url')) or ik
            parsed.append((im.get('position', len(parsed)), ik, tk))
        if not parsed:
            skipped += 1
            continue
        parsed.sort(key=lambda x: x[0])
        added_at = _parse_dt(rec.get('date_added'))
        cur.execute(
            'INSERT INTO items (id, name, category, image_key, thumb_key, date_added) '
            'VALUES (%s, %s, %s, %s, %s, %s)',
            (iid, rec.get('name') or '', rec.get('category') or '',
             parsed[0][1], parsed[0][2], added_at)
        )
        for pos, ik, tk in parsed:
            cur.execute(
                'INSERT INTO item_images (id, item_id, image_key, thumb_key, position, date_added) '
                'VALUES (%s, %s, %s, %s, %s, %s)',
                (uuid.uuid4().hex, iid, ik, tk, pos, added_at)
            )
        have.add(iid)
        added += 1
        photos += len(parsed)
    conn.commit()
    cur.close()
    conn.close()
    return added, photos, skipped


def import_from_zip(raw):
    """Full restore from a ZIP export: re-upload every photo file to R2 and
    recreate the rows."""
    zf = zipfile.ZipFile(io.BytesIO(raw))
    csv_name = next((n for n in zf.namelist() if n.lower().endswith('.csv')), None)
    if not csv_name:
        raise ValueError('no catalog.csv inside the zip')
    rows = list(csv.DictReader(io.StringIO(zf.read(csv_name).decode('utf-8-sig'))))
    added = photos = skipped = 0
    s3 = get_r2_client()
    conn = get_db()
    cur = conn.cursor()
    have = _existing_ids(cur)
    for row in rows:
        iid = (row.get('id') or '').strip()
        if not iid or iid in have:
            skipped += 1
            continue
        files = [x.strip() for x in (row.get('photo_files') or '').split('|') if x.strip()]
        saved = []
        for fname in files:
            try:
                body = zf.read(f'images/{fname}')
            except KeyError:
                continue
            ext = fname.rsplit('.', 1)[-1].lower() if '.' in fname else 'jpg'
            uid = uuid.uuid4().hex
            image_key, thumb_key = f'items/{uid}.{ext}', f'items/thumbs/{uid}.jpg'
            s3.put_object(Bucket=R2_BUCKET_NAME, Key=image_key, Body=body,
                          ContentType=f'image/{ext}')
            try:
                s3.put_object(Bucket=R2_BUCKET_NAME, Key=thumb_key,
                              Body=make_thumbnail_bytes(body), ContentType='image/jpeg')
            except Exception:
                thumb_key = image_key
            saved.append((image_key, thumb_key))
        if not saved:
            skipped += 1
            continue
        added_at = _parse_dt(row.get('date_added'))
        cur.execute(
            'INSERT INTO items (id, name, category, image_key, thumb_key, date_added) '
            'VALUES (%s, %s, %s, %s, %s, %s)',
            (iid, row.get('name') or '', row.get('category') or '',
             saved[0][0], saved[0][1], added_at)
        )
        for pos, (ik, tk) in enumerate(saved):
            cur.execute(
                'INSERT INTO item_images (id, item_id, image_key, thumb_key, position, date_added) '
                'VALUES (%s, %s, %s, %s, %s, %s)',
                (uuid.uuid4().hex, iid, ik, tk, pos, added_at)
            )
        have.add(iid)
        added += 1
        photos += len(saved)
    conn.commit()
    cur.close()
    conn.close()
    return added, photos, skipped


def _flash_import_result(added, photos, skipped):
    flash(t('import_done').format(added=added, photos=photos, skipped=skipped), 'success')


@app.route('/import', methods=['GET', 'POST'])
@login_required
def import_data():
    if request.method == 'POST':
        f = request.files.get('file')
        if not f or not f.filename:
            flash(t('import_no_file'), 'error')
            return redirect(url_for('import_data'))
        name = f.filename.lower()
        try:
            if name.endswith('.zip'):
                result = import_from_zip(f.read())
            elif name.endswith('.json'):
                result = import_from_json(f.read())
            else:
                flash(t('import_bad_type'), 'error')
                return redirect(url_for('import_data'))
        except Exception as exc:
            flash(f"{t('import_failed')} ({exc})", 'error')
            return redirect(url_for('import_data'))
        _flash_import_result(*result)
        return redirect(url_for('index'))
    return render_template('import.html')


@app.route('/import/latest', methods=['POST'])
@login_required
def import_latest():
    try:
        raw = r2_get_bytes('backups/latest.json')
    except Exception:
        flash(t('import_nolatest'), 'error')
        return redirect(url_for('import_data'))
    try:
        result = import_from_json(raw)
    except Exception as exc:
        flash(f"{t('import_failed')} ({exc})", 'error')
        return redirect(url_for('import_data'))
    _flash_import_result(*result)
    return redirect(url_for('index'))


# ---------------------------------------------------------------------------
# Routes — admin only
# ---------------------------------------------------------------------------
@app.route('/manage')
@login_required
def manage():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) FROM items')
    total = cur.fetchone()[0]
    cur.close()
    conn.close()
    return render_template('manage.html', total=total)


@app.route('/upload', methods=['GET', 'POST'])
@login_required
def upload():
    if request.method == 'POST':
        item_id = request.form.get('item_id', '').strip()
        name = request.form.get('name', '').strip()
        category = request.form.get('category', '').strip()
        files = [f for f in request.files.getlist('images') if f and f.filename]

        if not item_id:
            flash(t('id_required'), 'error')
            return redirect(url_for('upload'))
        if not files:
            flash(t('image_required'), 'error')
            return redirect(url_for('upload'))
        if any(not allowed_file(f.filename) for f in files):
            flash(t('bad_filetype'), 'error')
            return redirect(url_for('upload'))

        if get_item(item_id):
            flash(t('id_exists'), 'error')
            return redirect(url_for('upload'))

        saved = [save_image(f) for f in files]
        now = datetime.now(timezone.utc)
        cover_image, cover_thumb = saved[0]

        conn = get_db()
        cur = conn.cursor()
        cur.execute(
            'INSERT INTO items (id, name, category, image_key, thumb_key, date_added) VALUES (%s, %s, %s, %s, %s, %s)',
            (item_id, name, category, cover_image, cover_thumb, now)
        )
        for position, (image_key, thumb_key) in enumerate(saved):
            cur.execute(
                'INSERT INTO item_images (id, item_id, image_key, thumb_key, position, date_added) VALUES (%s, %s, %s, %s, %s, %s)',
                (uuid.uuid4().hex, item_id, image_key, thumb_key, position, now)
            )
        conn.commit()
        cur.close()
        conn.close()
        flash(t('item_added'), 'success')
        return redirect(url_for('upload'))

    return render_template('upload.html')


@app.route('/item/<item_id>/edit', methods=['GET', 'POST'])
@login_required
def edit_item(item_id):
    item = get_item(item_id)
    if not item:
        flash(t('item_not_found'), 'error')
        return redirect(url_for('index'))

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        category = request.form.get('category', '').strip()
        remove_ids = set(request.form.getlist('remove'))
        new_files = [f for f in request.files.getlist('images') if f and f.filename]

        if any(not allowed_file(f.filename) for f in new_files):
            flash(t('bad_filetype'), 'error')
            return redirect(url_for('edit_item', item_id=item_id))

        current = get_item_images(item_id)
        keeping = [img for img in current if img['id'] not in remove_ids]
        removing = [img for img in current if img['id'] in remove_ids]

        if not keeping and not new_files:
            flash(t('image_required'), 'error')
            return redirect(url_for('edit_item', item_id=item_id))

        saved = [save_image(f) for f in new_files]
        now = datetime.now(timezone.utc)
        next_position = max([img['position'] for img in keeping], default=-1) + 1

        conn = get_db()
        cur = conn.cursor()
        for img in removing:
            cur.execute('DELETE FROM item_images WHERE id = %s', (img['id'],))
        for offset, (image_key, thumb_key) in enumerate(saved):
            cur.execute(
                'INSERT INTO item_images (id, item_id, image_key, thumb_key, position, date_added) VALUES (%s, %s, %s, %s, %s, %s)',
                (uuid.uuid4().hex, item_id, image_key, thumb_key, next_position + offset, now)
            )
        cur.execute(
            'SELECT image_key, thumb_key FROM item_images WHERE item_id = %s ORDER BY position, date_added LIMIT 1',
            (item_id,)
        )
        cover = cur.fetchone()
        cur.execute(
            'UPDATE items SET name = %s, category = %s, image_key = %s, thumb_key = %s WHERE id = %s',
            (name, category, cover[0], cover[1], item_id)
        )
        conn.commit()
        cur.close()
        conn.close()

        delete_all_item_images(removing)
        flash(t('item_updated'), 'success')
        return redirect(url_for('item_detail', item_id=item_id))

    return render_template('edit.html', item=item, images=get_item_images(item_id))


@app.route('/item/<item_id>/delete', methods=['POST'])
@login_required
def delete_item(item_id):
    item = get_item(item_id)
    if item:
        images = get_item_images(item_id)
        conn = get_db()
        cur = conn.cursor()
        cur.execute('DELETE FROM item_images WHERE item_id = %s', (item_id,))
        cur.execute('DELETE FROM items WHERE id = %s', (item_id,))
        conn.commit()
        cur.close()
        conn.close()
        delete_all_item_images(images)
    return redirect(url_for('index'))


# Make sure the schema exists on startup — safe to run repeatedly, including
# under gunicorn where __main__ below never executes.
try:
    init_db()
except Exception as exc:  # best effort; real errors surface on the first request
    print(f'[init_db] schema setup skipped: {exc}')


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
